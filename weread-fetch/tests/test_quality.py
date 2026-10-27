"""离线回归：真实长段落的高亮边界、alt XHTML 脚注、仿宋原样保留、验收拒绝坏件。"""
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from lib import docxbuild
from lib.cli import clean_author, clean_title, title_needs_review
from lib.docxcheck import audit


class QualityTests(unittest.TestCase):
    def test_book_and_rejected_highlight(self):
        work = Path(tempfile.mkdtemp(prefix="weread-quality-test-"))
        (work / "xhtml").mkdir()
        (work / "css").mkdir()
        (work / "book.json").write_text(json.dumps({"bookId": "test", "title": "质量测试",
            "author": "测试", "chapters": [{"chapterUid": 1, "title": "测试章节", "level": 1}]}))
        (work / "progress.json").write_text(json.dumps({"1": {"ok": True, "file": "xhtml/1.xhtml"}}))
        long_text = "正常文字" * 100 + "\ue544" + "后续正常文字" * 100
        (work / "xhtml/1.xhtml").write_text(
            '<html><body><h1>测试章节</h1><p class="preface-text">' + long_text + '</p>'
            '<p>正文<img class="qqreader-footnote" alt="&lt;p&gt;脚注&lt;b&gt;强调&lt;/b&gt;'
            '\ue583&lt;br/&gt;下一行&lt;/p&gt;"/></p>'
            '<p>扩展私用区\U000f0001和\U00100001保留。<a href="#n1"><sup>1</sup></a></p>'
            '<p class="note" id="n1"><a href="#n1-back">1</a> &lt;p class="zy27"&gt;'
            '转义注文\ue583末尾&lt;/p&gt;</p></body></html>')
        (work / "css/1.css").write_text('.preface-text {font-family: "STFangsong", "Fangsong", "仿宋";}')
        out = work / "quality.docx"
        with patch('lib.pua.resolve', return_value={}), patch('lib.pua.load_manual', return_value={}):
            stats = docxbuild.build_book(work, out, docxbuild.Options(toc=False))
        report = stats['audit']
        self.assertEqual(report['errors'], [])
        self.assertEqual(report['highlighted_characters'], 5)
        self.assertEqual(report['footnotes'], 2)
        d = Document(out)
        para = next(p for p in d.paragraphs if p.text == long_text)
        self.assertTrue(all(r._r.rPr.rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia') == '仿宋' for r in para.runs))
        self.assertEqual(''.join(r.text for r in para.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW), '\ue544')
        from zipfile import ZipFile
        from xml.etree import ElementTree as ET
        with ZipFile(out) as z:
            root = ET.fromstring(z.read('word/footnotes.xml'))
            text = ''.join(root.itertext())
            self.assertIn('脚注强调\ue583下一行', text)
            self.assertNotIn('<p>', text)
            self.assertNotIn('<p class', text)
            self.assertIn('转义注文\ue583末尾', text)
            self.assertTrue(root.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br'))
        # 人为制造旧 bug；检测器必须失败，不能仅验证正常样本。
        para.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
        bad = work / 'bad.docx'
        d.save(bad)
        self.assertTrue(any('高亮超出' in e for e in audit(bad)['errors']))
        before = out.read_bytes()
        with patch('lib.pua.resolve', return_value={}), patch('lib.pua.load_manual', return_value={}), \
             patch('lib.docxcheck.audit', return_value={'errors': ['测试失败'], 'warnings': []}):
            with self.assertRaisesRegex(RuntimeError, '验收失败'):
                docxbuild.build_book(work, out, docxbuild.Options(toc=False))
        self.assertEqual(out.read_bytes(), before, '失败不能覆盖原成品')
        self.assertEqual(list(out.parent.glob('.weread-pending-*')), [], '失败不能在成品目录留隐藏候选件')
        self.assertTrue((work / 'docx-rejected.docx').exists(), '候选件挪回任务目录留查')


def make_book(chapters: list[dict], css: str) -> Path:
    """人造任务缓存。chapters 每项含 uid、title、level、files、xhtml。"""
    work = Path(tempfile.mkdtemp(prefix="weread-structure-test-"))
    (work / "xhtml").mkdir()
    (work / "css").mkdir()
    (work / "book.json").write_text(json.dumps({"bookId": "test", "title": "结构测试", "author": "测试",
        "chapters": [{"chapterUid": c["uid"], "chapterIdx": i + 1, "title": c["title"],
                      "level": c.get("level", 1), "files": c.get("files", [])}
                     for i, c in enumerate(chapters)]}, ensure_ascii=False), encoding="utf-8")
    progress = {}
    for c in chapters:
        stem = f"{c['uid']:04d}"
        (work / "xhtml" / f"{stem}.xhtml").write_text(c["xhtml"], encoding="utf-8")
        (work / "css" / f"{stem}.css").write_text(css, encoding="utf-8")
        progress[str(c["uid"])] = {"ok": True, "file": f"xhtml/{stem}.xhtml"}
    (work / "progress.json").write_text(json.dumps(progress), encoding="utf-8")
    return work


def build(work: Path) -> tuple[dict, Path]:
    out = work / "out.docx"
    with patch('lib.pua.resolve', return_value={}), patch('lib.pua.load_manual', return_value={}):
        stats = docxbuild.build_book(work, out, docxbuild.Options(toc=False))
    return stats, out


W_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
KINDLE_CSS = ('.content {font-family: "汉仪旗黑50S","HYQiHei-50s";}\n'
              '.quotation {font-family: "汉仪楷体","ETrump KaiTi";}\n'
              '.super {font-size: 0.7em; vertical-align: super;}\n'
              '.sub {font-size: 0.7em; vertical-align: sub;}\n')


class StructureTests(unittest.TestCase):
    """《亲密关系（第8版）》暴露的四类问题：拆行的章标题、多文件章的黑体、跨文件的注、验收盲区。"""

    LONG = "这是章首导语的正文，用来判断整本正文是不是黑体。"

    def split_title_book(self) -> Path:
        body = f'<p class="content">{self.LONG}</p>'
        kai = ''.join(f'<p class="quotation">• 第{i}条要点，楷体排的列表内容，足够长以计入长段落。</p>'
                      for i in range(6))
        return make_book([
            # 标题页 + 导语页两个文件拼成一章；四字以下的「第1章」认不上、章名能认上
            {"uid": 1, "title": "第1章 人际关系的构成", "files": ["Text/S1.xhtml", "Text/S1_0.xhtml"],
             "xhtml": '<html><body><h1 class="label">第1章</h1><p class="deco"></p>'
                      '<h1 class="name">人际关系的构成</h1></body></html>\n'
                      f'<html><body>{body}</body></html>'},
            # 「第10章」四个字会被旧判据当成整条标题
            {"uid": 2, "title": "第10章 压力与紧张", "files": ["Text/S10.xhtml"],
             "xhtml": f'<html><body><h1>第10章</h1><h1>压力与紧张</h1>{body}</body></html>'},
            # 两截都不足四个字；本章黑体只占 1/7，按章判不会转宋体，得靠全书判
            {"uid": 3, "title": "第3章 吸引力", "files": ["Text/S3.xhtml"],
             "xhtml": f'<html><body><h1>第3章</h1><h1>吸引力</h1>{body}{kai}</body></html>'},
        ], KINDLE_CSS)

    def test_split_title_and_heiti(self):
        stats, out = build(self.split_title_book())
        d = Document(out)
        heads = [(p.style.name, p.text) for p in d.paragraphs if p.style.name.startswith("WR 标题")]
        self.assertEqual(heads, [("WR 标题一", "第1章 人际关系的构成"), ("WR 标题一", "第10章 压力与紧张"),
                                 ("WR 标题一", "第3章 吸引力")])
        body = [p for p in d.paragraphs if p.text == self.LONG]
        self.assertEqual(len(body), 3)
        for p in body:
            fonts = {r._r.rPr.rFonts.get(W_NS + 'eastAsia') for r in p.runs}
            self.assertNotIn("黑体", fonts)
        warnings = "".join(stats["audit"]["warnings"])
        for word in ("对不上", "黑体", "重复"):
            self.assertNotIn(word, warnings)

    def test_audit_catches_old_defects(self):
        from lib.docxcheck import heading_mismatches
        _, out = build(self.split_title_book())
        d = Document(out)
        paras = d.paragraphs
        head = next(p for p in paras if p.text == "第3章 吸引力")
        nxt = paras[paras.index(head) + 1]
        nxt.insert_paragraph_before("第3章", style="WR 标题二")
        nxt.insert_paragraph_before("吸引力", style="WR 标题二")
        for run in next(p for p in paras if p.text == self.LONG).runs:
            run._r.get_or_add_rPr().get_or_add_rFonts().set(W_NS + 'eastAsia', "黑体")
        first = next(p for p in paras if p.text == "第1章 人际关系的构成")
        for run in first.runs:
            run.text = ""
        first.runs[0].text = "人际关系的构成"
        bad = out.with_name("bad.docx")
        d.save(bad)
        warnings = "".join(audit(bad)["warnings"])
        self.assertIn("整段仍是黑体", warnings)
        self.assertIn("第3章 吸引力 → 第3章 + 吸引力", warnings)
        expected = [("第1章 人际关系的构成", 1), ("第10章 压力与紧张", 1), ("第3章 吸引力", 1)]
        self.assertEqual(len(heading_mismatches(bad, expected)), 1)
        self.assertEqual(heading_mismatches(out, expected), [])

    def test_cross_file_notes(self):
        work = make_book([
            {"uid": 1, "title": "正文", "files": ["Text/A.xhtml"],
             "xhtml": '<html><body><h1>正文</h1><p>正文一句<span class="super">'
                      '<a href="../Text/B.xhtml#note_1" id="noteBack_1">[1]</a></span>后文。</p>'
                      '<p>水是H<span class="sub">2</span>O。</p></body></html>'},
            {"uid": 2, "title": "本章小结", "files": ["Text/B.xhtml"],
             "xhtml": '<html><body><h1>本章小结</h1><p>小结正文。</p><hr/>'
                      '<p class="quotation"><a href="../Text/A.xhtml#noteBack_1" id="note_1">[1]</a>'
                      '注文<img class="qqreader-footnote" src="../Images/note.png" alt="U.S. Census Bureau"/>'
                      '结束。</p></body></html>'},
        ], KINDLE_CSS)
        stats, out = build(work)
        self.assertEqual(stats["cross_file_notes"], 1)
        self.assertEqual(stats["audit"]["footnotes"], 1)
        from zipfile import ZipFile
        from xml.etree import ElementTree as ET
        with ZipFile(out) as z:
            notes = ''.join(ET.fromstring(z.read('word/footnotes.xml')).itertext())
        self.assertIn("注文（U.S. Census Bureau）结束。", notes)
        self.assertNotIn("[1]", notes)
        d = Document(out)
        text = "\n".join(p.text for p in d.paragraphs)
        self.assertNotIn("[1]", text)
        self.assertNotIn("注文", text)
        self.assertNotIn("WR 分隔", [p.style.name for p in d.paragraphs], "注文挪走后分隔线不该留下")
        sub = next(r for p in d.paragraphs if p.text == "水是H2O。" for r in p.runs if r.text == "2")
        self.assertTrue(sub.font.subscript)
        self.assertEqual(sub.font.size.pt, 11.0, "CSS 的 0.7em 不能叠在 Word 下标的缩小上")

    def test_notes_split_across_files_of_one_chapter(self):
        # 《波斯笔记》第四章：注号与注文在同一章的两个原文件里，锚点带文件名、注文没有 id
        work = make_book([
            {"uid": 1, "title": "第四章 驿道", "files": ["Text/C12.html", "Text/C13.html"],
             "xhtml": '<html><body><h1>第四章 驿道</h1><p>正文<a href="../Text/C13.html#ch1" '
                      'id="ch1-back"><sup>(1)</sup></a>。</p></body></html>\n'
                      '<html><body><p>更多正文。</p><hr/><p class="fnote">'
                      '<a href="../Text/C12.html#ch1-back">(1)</a>&nbsp;注文甲。</p></body></html>'},
        ], KINDLE_CSS)
        stats, out = build(work)
        self.assertEqual(stats["audit"]["footnotes"], 1)
        d = Document(out)
        self.assertNotIn("注文甲", "\n".join(p.text for p in d.paragraphs))
        self.assertNotIn("WR 分隔", [p.style.name for p in d.paragraphs])

    def test_images_word_cannot_read(self):
        """封面图床的 JPEG 没有 JFIF 头、正文里 Exif 段坏掉的 JPEG：都得进 Word，不能静默丢。"""
        import hashlib
        import io
        from PIL import Image
        from docx.image.image import Image as DocxImage
        buf = io.BytesIO()
        Image.new("RGB", (60, 80), (180, 40, 40)).save(buf, "JPEG")
        data = buf.getvalue()
        raw = data[:2] + data[4 + int.from_bytes(data[4:6], "big"):]      # 剥掉 JFIF 段
        app1 = b"Exif\x00\x00II*\x00" + (9999).to_bytes(4, "little")    # IFD 偏移指到文件外
        bad_exif = raw[:2] + b"\xff\xe1" + (len(app1) + 2).to_bytes(2, "big") + app1 + raw[2:]
        for blob in (raw, bad_exif):
            with self.assertRaises(Exception, msg="夹具前提：python-docx 原样认不出"):
                DocxImage.from_blob(blob)
        work = make_book([
            {"uid": 1, "title": "封面", "xhtml": "<html><body><p>版权页内容</p></body></html>"},
            {"uid": 2, "title": "第一章", "xhtml": '<html><body><h1>第一章</h1><p>正文。</p>'
             '<div><img alt="" class="qqreader-fullimg" src="https://example.com/raw.jpg"/></div>'
             '<div><img alt="" class="qqreader-fullimg" src="https://example.com/exif.jpg"/></div>'
             '</body></html>'},
        ], "")
        meta = json.loads((work / "book.json").read_text(encoding="utf-8"))
        meta["cover"] = "https://example.com/cover.jpg"
        (work / "book.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
        (work / "images").mkdir()
        cached = {}
        for url, blob in ((meta["cover"], raw), ("https://example.com/raw.jpg", raw),
                          ("https://example.com/exif.jpg", bad_exif)):
            f = work / "images" / (hashlib.md5(url.encode()).hexdigest()[:16] + ".jpg")
            f.write_bytes(blob)
            cached[f] = blob
        stats, out = build(work)
        self.assertEqual(stats["cover"], "嵌入")
        self.assertEqual(stats["images"], 2)
        self.assertEqual(stats["image_embed_failures"], 0)
        self.assertEqual(stats["audit"]["warnings"], [])
        d = Document(out)
        self.assertEqual(d.paragraphs[2].style.name, "WR 图")
        self.assertTrue(d.paragraphs[2]._p.findall(".//" + W_NS + "drawing"), "扉页第三段是封面图")
        text = "\n".join(p.text for p in d.paragraphs)
        self.assertNotIn("版权页内容", text, "封面进了扉页，封面章照旧跳过")
        self.assertNotIn("图片无法嵌入", text)
        for f, blob in cached.items():
            self.assertEqual(f.read_bytes(), blob, "缓存里的原图不能被改")
        # 人为造出老毛病：空的图段落，验收必须报
        d.add_paragraph(style="WR 图")
        bad = work / "bad.docx"
        d.save(bad)
        self.assertTrue(any("空的图段落" in w for w in audit(bad)["warnings"]))

    def test_toc_follows_rendered_headings(self):
        """书前目录照正文实际排出的标题：接口不单列的同文件小节（《失控》的 x.1）也得在。"""
        self.assertFalse(docxbuild.Options().toc, "默认不排书前目录，靠侧栏大纲")
        work = make_book([
            {"uid": 1, "title": "第一章 人造与天生",
             "xhtml": '<html><body><h1>第一章 人造与天生</h1><h2>1.1 新生物文明</h2><p>正文一。</p>'
                      '<h3>一个小标题</h3><p>正文。</p></body></html>'},
            {"uid": 2, "title": "1.2 生物逻辑的胜利", "level": 2,
             "xhtml": '<html><body><h2>1.2 生物逻辑的胜利</h2><p>正文二。</p></body></html>'},
        ], "")
        out = work / "toc.docx"
        with patch('lib.pua.resolve', return_value={}), patch('lib.pua.load_manual', return_value={}):
            stats = docxbuild.build_book(work, out, docxbuild.Options(toc=True))
        d = Document(out)
        toc = [(p.style.name, p.text) for p in d.paragraphs if p.style.name.startswith("WR 目录")]
        self.assertEqual(toc, [("WR 目录1", "第一章 人造与天生"), ("WR 目录2", "1.1 新生物文明"),
                               ("WR 目录2", "1.2 生物逻辑的胜利")], "深度跟章节表走，小标题不进目录")
        self.assertEqual(stats["toc_entries"], 3)
        self.assertEqual(stats["audit"]["warnings"], [])
        names = [p.style.name for p in d.paragraphs]
        self.assertLess(names.index("WR 目录1"), names.index("WR 标题二"), "目录在正文之前")


class TitleTests(unittest.TestCase):
    """题名清理：营销话术该删，版次册数丛书编号一个字不能动。"""

    def test_marketing_paren_removed(self):
        self.assertEqual(clean_title('智能：AI时代的商业、组织与战略的本质（曾鸣教授全新商业思考）'),
                         '智能：AI时代的商业、组织与战略的本质')
        self.assertEqual(clean_title('某某传（全新解读）'), '某某传')
        self.assertEqual(clean_title('某某传（官方旗舰店独家首发）'), '某某传')

    def test_marketing_bracket_removed(self):
        self.assertEqual(clean_title('某某集【中华书局出品】'), '某某集')
        self.assertEqual(clean_title('某某集【赠精美书签】'), '某某集')

    def test_biblio_paren_kept(self):
        for title in ('三家《诗》辑佚史（增订版）', '某某（第8版）', '某某（修订本）',
                      '某某（影印本）', '某某【译丛067】'):
            self.assertEqual(clean_title(title), title, title)

    def test_cleaned_title_needs_no_review(self):
        self.assertEqual(title_needs_review(clean_title('某某（曾鸣教授全新商业思考）')), '')
        # 平台造了新词、没删干净时，得把可疑括号亮出来
        self.assertIn('（全新商业思考）', title_needs_review('某某（全新商业思考）'))

    def test_biblio_bookname_field_cleaned(self):
        from bs4 import BeautifulSoup
        soup = BeautifulSoup('<html><body><p>书名：某某传（全新解读）</p>'
                             '<p>出版社：某某社</p><p>ISBN：123</p></body></html>',
                             'html.parser')
        n = docxbuild.strip_marketing_from_biblio(soup.body)
        self.assertEqual(n, 1)
        self.assertEqual(soup.body.find_all('p')[0].get_text(strip=True), '书名：某某传')
        self.assertEqual(soup.body.find_all('p')[1].get_text(strip=True), '出版社：某某社')
        self.assertEqual(soup.body.find_all('p')[2].get_text(strip=True), 'ISBN：123')

    def test_biblio_bookname_without_marketing_untouched(self):
        from bs4 import BeautifulSoup
        soup = BeautifulSoup('<html><body><p>书名：某某传（增订版）</p></body></html>',
                             'html.parser')
        self.assertEqual(docxbuild.strip_marketing_from_biblio(soup.body), 0)
        self.assertEqual(soup.p.get_text(strip=True), '书名：某某传（增订版）')

    def test_author_role_kept_and_multi_author_untouched(self):
        self.assertEqual(clean_author('马昕著'), '马昕')
        self.assertEqual(clean_author('汤用彤 汤一介等导读'), '汤用彤 汤一介等导读')

    def test_volume_count_removed(self):
        # 几册已合成一本电子书，册数不是题名（《失控》的「（上、下）」曾一路带进文件名）
        cases = {
            '失控：全人类的最终命运和结局（上、下）': '失控：全人类的最终命运和结局',
            '某某（上下）': '某某', '某某（上中下）': '某某', '某某（上、中、下）': '某某',
            '某某（上、下册）': '某某', '某某（上下卷）': '某某', '某某（全两册）': '某某',
            '某某（3册）': '某某', '某某（套装共3册）': '某某', '某某（套装上下册）': '某某',
            '三家《诗》辑佚史（精）全二册': '三家《诗》辑佚史',
        }
        for raw, want in cases.items():
            self.assertEqual(clean_title(raw), want, raw)
            self.assertEqual(title_needs_review(clean_title(raw)), '', raw)

    def test_single_volume_kept(self):
        # 只收一册的书，删了会和另一册撞成同一个文件名
        for title in ('某某（上）', '某某（下册）', '某某（第一册）', '某某第三册',
                      '第二性（合卷本）', '余音：学术史随笔选（1992—2015）'):
            self.assertEqual(clean_title(title), title, title)
            self.assertEqual(title_needs_review(title), '', title)

    def test_uncertain_volume_flagged_not_guessed(self):
        # 区间是不是全套、括注里夹着的册数删了会剩半个括号：都不动，亮出来交给人
        for title, left in (('某某全集（1-3册）', '（1-3册）'),
                            ('某某（插图珍藏版·全三册）', '（插图珍藏版·全三册）')):
            self.assertEqual(clean_title(title), title)
            self.assertIn(left, title_needs_review(title))

    def test_biblio_bookname_volume_cleaned(self):
        from bs4 import BeautifulSoup
        soup = BeautifulSoup('<html><body><p>书名：失控：全人类的最终命运和结局（上、下）</p></body></html>',
                             'html.parser')
        self.assertEqual(docxbuild.strip_marketing_from_biblio(soup.body), 1)
        self.assertEqual(soup.p.get_text(strip=True), '书名：失控：全人类的最终命运和结局')


if __name__ == '__main__':
    unittest.main()
