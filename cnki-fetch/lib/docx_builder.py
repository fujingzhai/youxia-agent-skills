# -*- coding: utf-8 -*-
"""把 CNKI AI 阅读页的 DOM 排成 DOCX。

版式（我们统一的部分）：
    来源 / DOI / 题名 / 作者·单位 / 摘要 / 关键词 / 正文
    全宋体；正文小四 1.5 倍行距首行缩进 2 字；表题图题、参考文献五号；脚注小五单倍；A4，上下 2.54cm 左右 3.17cm
    标题 小二 18 → 三号 16 → 四号 14 → 小四 12 → 小四 12，前两级居中（四级不小于正文）
    脚注为真·页下脚注，全篇连续编号；基金作题名的星号注、作者简介作作者行的星号注，都不占编号
    英文题名、英文作者、英文摘要一律不收；来源行随正文繁简

照原刊不动的部分：章节措辞与层级、注释引书格式、繁简、摘要关键词的有无。
"""
import io, re, os, tempfile
from bs4 import BeautifulSoup
from PIL import Image
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import footnote as FN
import quality as QA
import records as REC

BUILD_VERSION = 3

# 使用 macOS/LibreOffice 都能解析的 CJK 宋体族名；本机“宋体”别名在
# headless LibreOffice 中可能退化为无字形字体，导致验收图出现方框。
FONT = "STSong"
SZ_BODY = Pt(12)       # 正文：小四
SZ_AUX = Pt(10.5)      # 表题图题、参考文献：五号，比正文小一号
SZ_NOTE = Pt(9)        # 脚注：小五
SZ_META = Pt(10.5)     # 来源 / DOI / 作者：五号
SZ_H0 = Pt(18)         # 文章标题：小二
SZ_H1, SZ_H2, SZ_H3, SZ_H4 = Pt(16), Pt(14), Pt(12), Pt(12)     # 三号 / 四号 / 小四 / 小四

_TRAD = set("為體學國關書禮這對發說時實樣們來經義傳認種舊語當萬與從業點術產處")
_SIMP = set("为体学国关书礼这对发说时实样们来经义传认种旧语当万与从业点术产处")

try:
    from opencc import OpenCC
    _S2T, _T2S = OpenCC('s2t'), OpenCC('t2s')
except Exception:          # 没装 opencc 就照抄原文，不至于整篇生成不出来
    _S2T = _T2S = None


def _cc(text, trad):
    """把知网给的简体元数据转成与正文一致的用字"""
    conv = _S2T if trad else _T2S
    return conv.convert(text) if conv else text


def is_trad(text):
    """按正文用字判断繁简，用来决定字段名写「关键词」还是「關鍵詞」。"""
    t = sum(text.count(c) for c in _TRAD)
    s = sum(text.count(c) for c in _SIMP)
    return t >= s


def preclean(soup, trad):
    """只改字段标签，不再清空段落（其中可能有图片、注释和非 span 文字）。"""
    labels = {"摘要:": "摘要：", "关键词:": "關鍵詞：" if trad else "关键词：",
              "基金:": "基金：", "作者:": "作者："}
    for p in soup.select('div.ChapterWrap.Content p'):
        b = p.find('span', class_='b')
        if not b:
            continue
        label = b.get_text().strip()
        if label not in labels:
            continue
        b.string = labels[label]
        if label == '关键词:':
            vals = [sp for sp in b.find_next_siblings('span') if sp.get_text(strip=True)]
            for sp in vals[:-1]:
                if not sp.get_text().rstrip().endswith((';', '；', '、', ',' , '，')):
                    sp.insert_after('；')


def people(soup):
    """ul.Chapter-people：中文作者 / 单位 / 英文作者"""
    return [re.sub(r'\s+', ' ', ul.get_text(' ', strip=True)) for ul in soup.select('ul.Chapter-people')]


def parse(soup, refs=None):
    """→ (blocks, notes, img_srcs)。block = (标签名, 混排节点)。
    节点 = ('t', 文本) | ('n', 脚注序号) | ('i', 图片在正文中的序号)"""
    # 全部图片预先按原始 DOM 编号，兼容旧缓存；表格/精简模式不得改变序号。
    # 只有 attachment URL 有对应的浏览器响应/缓存字节；其它 XML img 节点
    # （例如公式占位节点）没有可取的二进制，不能虚构成已抓到的图片。
    images = soup.select('div.ChapterWrap.Content img[src*="attachment"]')
    srcs = [im.get('src', '') for im in images]
    image_ids = {id(im): i for i, im in enumerate(images)}
    notes = []

    def inline(el):
        if refs:
            mark_refs(el)
        nodes = []
        def visit(d):
            name = getattr(d, 'name', None)
            if isinstance(d, str):
                if d.strip():
                    txt = _degap(re.sub(r'[\r\n\t]+', ' ', str(d)))
                    nodes.extend(split_refs(txt, refs, notes) if refs else [('t', txt)])
                elif nodes and ' ' in d:
                    nodes.append(('t', ' '))
            elif name in ('script', 'style'):
                return
            elif name == 'sup' and 'footNote' in (d.get('class') or []):
                notes.append(_degap(re.sub(r'^\(\d+\)\s*', '', d.get('footnotetitle', '').strip())))
                nodes.append(('n', len(notes)))
            elif name == 'img':
                if id(d) in image_ids:
                    nodes.append(('i', image_ids[id(d)]))
                else:
                    ref = d.get('src') or d.get('href') or d.get('alt') or '无源地址'
                    nodes.append(('t', f'［图片：{ref}］'))
            elif name == 'br':
                nodes.append(('t', '\n'))
            elif name in ('sup', 'sub') and not d.find(['img', 'sup', 'sub']):
                nodes.append((name, d.get_text()))
            else:
                for ch in d.children:
                    visit(ch)
                if name in ('p', 'div', 'li') and d is not el:
                    nodes.append(('t', '\n'))
        visit(el)
        return nodes

    def table(el):
        rows = [r for r in el.find_all('tr') if r.find_parent('table') is el]
        occupied, cells, width = set(), [], 0
        for r, row in enumerate(rows):
            c = 0
            for td in row.find_all(['td', 'th'], recursive=False):
                while (r, c) in occupied:
                    c += 1
                rs = int(td.get('rowspan', 1) or 1)
                rs = len(rows) - r if rs == 0 else min(rs, len(rows) - r)
                cs = max(1, int(td.get('colspan', 1) or 1))
                if rs < 1 or cs > 100:
                    raise ValueError('表格跨行/跨列数据异常')
                cells.append({'r': r, 'c': c, 'rs': rs, 'cs': cs,
                              'header': td.name == 'th', 'blocks': walk(td)})
                for rr in range(r, r + rs):
                    for cc in range(c, c + cs):
                        if (rr, cc) in occupied:
                            raise ValueError('表格合并范围相互重叠')
                        occupied.add((rr, cc))
                c += cs; width = max(width, c)
        return {'rows': len(rows), 'cols': width, 'cells': cells}

    def walk(root):
        blocks, pending = [], []
        def flush():
            if pending:
                wrapper = BeautifulSoup('<span></span>', 'html.parser').span
                for ch in pending:
                    wrapper.append(ch)
                nodes = inline(wrapper)
                if nodes:
                    blocks.append(('p', nodes))
                pending.clear()
        # 不跨根扫描，避免表格格内段落再次出现在表格外。
        for el in list(root.children):
            name = getattr(el, 'name', None)
            if name in ('script', 'style'):
                continue
            if name in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p', 'caption'):
                flush(); nodes = inline(el)
                if nodes:
                    blocks.append((name, nodes))
            elif name == 'table':
                flush()
                caption = el.find('caption', recursive=False)
                if caption:
                    blocks.append(('caption', inline(caption)))
                blocks.append(('table', table(el)))
            elif name == 'img':
                flush(); blocks.append(('fig', inline(el)))
            elif name in ('div', 'section', 'article', 'ul', 'ol', 'li'):
                flush(); blocks.extend(walk(el))
            else:
                pending.append(el)
        flush()
        return blocks

    blocks = []
    for div in soup.select('div.ChapterWrap.Content'):
        if not div.find_parent('div', class_='ChapterWrap'):
            blocks.extend(walk(div))
    return blocks, notes, srcs


def text_nodes(nodes):
    return ''.join(v for k, v in nodes if k in ('t', 'sup', 'sub'))


_CJK = r'\u3400-\u9fff\uf900-\ufaff\u3000-\u303f\uff01-\uff65\u2014\u2018-\u201f\u2026\u2027'
_GAP = re.compile(rf'(?<=[{_CJK}])[ \u00a0]+(?=[{_CJK}])')


def _degap(t):
    """去掉汉字与汉字之间的游离空格。

    这类空格是知网 XML 自带的（原刊排版的字距／换行转换残留），如「义俭辞 碎」「文 心雕龙」。
    范围含汉字与中文标点（含中文引号、破折号、省略号）；中英之间、数字两侧的空格一律保留。
    """
    return _GAP.sub('', t)


def _is_en(t):
    """首部里的英文题名/摘要——西文字母过半即判定"""
    letters = sum(1 for c in t if c.isascii() and c.isalpha())
    return len(t) > 30 and letters / max(1, len(t)) > 0.5


def _has_han(t):
    """有没有汉字。英文作者名太短，占比法判不出来"""
    return bool(re.search(r'[\u3400-\u9fff]', t))


def references(soup):
    """尾注体例的刊物：正文里的 [1][2] 是普通文本，条目单独放在 div.references-box。
    不在 ChapterWrap 里，得单独取，否则文档只剩注号没有注文。"""
    box = soup.select_one('div.references-box')
    if not box:
        return []
    return [_degap(li.get_text(' ', strip=True)) for li in box.select('li') if li.get_text(strip=True)]


# 尾注引用在 DOM 里是两个相邻的 sup：<sup><a type=reference>[1]</a></sup><sup>(P349)</sup>
# 先把这一对合成一个标记，再按标记切分——纯文本正则接不上这两个节点
MARK = "\ufff9"      # 别用 \u0000：残留时 python-docx 会拒收 NULL 字节
REF_CITE = re.compile(MARK + r'(\d+)\|([^\ufff9]*)' + MARK + r'|\[(\d+)\](?:\(([^)（）]{0,60})\))?')


def mark_refs(el):
    """只处理编号制尾注 <sup><a type=reference>[1]</a></sup>。

    同一个 a[type=reference] 还用于著者出版年制的文内夹注（<citation><a>丁四新，2006</a></citation>
    年，第43页)），那种本身就完整可读，必须原样保留——按数字粗暴匹配会把「2006」当成编号，
    整条夹注就没了。
    """
    for a in el.select('a[type="reference"]'):
        if getattr(a.parent, 'name', None) != 'sup':
            continue
        if getattr(a.parent, 'parent', None) is None:
            continue
        mm = re.fullmatch(r'\s*\[(\d+)\]\s*', a.get_text())
        if not mm:
            continue
        sup, n = a.parent, mm.group(1)
        page, nxt = '', sup.next_sibling
        while nxt is not None and isinstance(nxt, str) and not nxt.strip():
            nxt = nxt.next_sibling
        if getattr(nxt, 'name', None) == 'sup':
            t = nxt.get_text().strip()
            if re.match(r'^\([Pp]', t):
                page = t.strip('()'); nxt.extract()
        sup.replace_with(f"{MARK}{n}|{page}{MARK}")


def _page_phrase(raw):
    """(P349) → 第349页；(P58—60、P84) → 第58—60、84页"""
    if not raw:
        return ""
    t = re.sub(r'[Pp]', '', raw).strip(' ,，')
    return f"，第{t}页" if t else ""


def split_refs(text, refs, notes):
    """尾注体例：把正文里的 [1](P349) 就地换成脚注引用。

    [1] 指参考文献第 1 条，(P349) 是该处引用的页码。同一条文献在文中被引十几次是常事，
    每处都生成一条独立脚注，各带各的页码——这正是页下注体例该有的样子。
    注文沿用原刊的著录格式，不改写。
    """
    out, last = [], 0
    for m in REF_CITE.finditer(text):
        num, page = (m.group(1), m.group(2)) if m.group(1) else (m.group(3), m.group(4))
        n = int(num)
        if not (0 < n <= len(refs)):        # 编号不在表内，还原成可读文本，别留标记
            if m.start() > last:
                out.append(("t", text[last:m.start()]))
            out.append(("t", f"[{num}]" + (f"({page})" if page else "")))
            last = m.end()
            continue
        if m.start() > last:
            out.append(("t", text[last:m.start()]))
        src = re.sub(r'^\s*\[\d+\]\s*', '', refs[n - 1]).rstrip('。．. ')
        notes.append(src + _page_phrase(page) + "。")
        out.append(("n", len(notes)))
        last = m.end()
    if last < len(text):
        out.append(("t", text[last:]))
    if not out:
        out = [("t", text)]
    return [(k, v.replace(MARK, '') if k == "t" else v) for k, v in out]


def _cnfont(run, size, bold=False):
    run.font.name = FONT; run.font.size = size; run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    run.font.color.rgb = RGBColor(0, 0, 0)


def _newp(doc, *, size, spacing, indent_chars=0, align=None, before=0, after=0):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.line_spacing = spacing; pf.space_before = Pt(before); pf.space_after = Pt(after)
    if indent_chars:
        pf.first_line_indent = size * indent_chars
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def _emit(p, nodes, size, imgs):
    """文字 / 脚注引用 / 生僻字图片按原顺序落进同一段"""
    for kind, val in nodes:
        if kind == "t":
            if val:
                _cnfont(p.add_run(val), size)
        elif kind in ('sup', 'sub'):
            run = p.add_run(val); _cnfont(run, size)
            run.font.superscript = kind == 'sup'
            run.font.subscript = kind == 'sub'
        elif kind == "n":
            FN.add_ref(p, val + 2, sz_half=int(size.pt * 2))   # id=1、2 留给基金与作者简介
        elif kind == "i":
            data = imgs.get(val)
            try:
                im = Image.open(io.BytesIO(data))
                p.add_run().add_picture(io.BytesIO(data), height=size,
                                        width=Pt(size.pt * im.width / im.height))
            except Exception:
                _cnfont(p.add_run("▯"), size)   # 取不到时的占位。用 ▯ 不用 □，避免与原文表残缺的 □ 混淆


def _emit_table(doc, spec, imgs, width=Cm(14.66)):
    if not spec['rows'] or not spec['cols']:
        return
    table = doc.add_table(rows=spec['rows'], cols=spec['cols'])
    table.style = 'Table Grid'; table.autofit = False
    for col in table.columns:
        col.width = int(width / spec['cols'])
    for cell in spec['cells']:
        r, c, rs, cs = (cell[k] for k in ('r', 'c', 'rs', 'cs'))
        target = table.cell(r, c)
        if rs > 1 or cs > 1:
            target = target.merge(table.cell(r + rs - 1, c + cs - 1))
        target.width = int(width * cs / spec['cols'])
        first = True
        for kind, nodes in cell['blocks']:
            if kind == 'table':
                _emit_table(target, nodes, imgs, int(width * cs / spec['cols']))
                continue
            p = target.paragraphs[0] if first else target.add_paragraph()
            first = False
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.first_line_indent = Pt(0)
            _emit(p, nodes, SZ_NOTE, imgs)
            if cell['header'] or kind == 'caption':
                for run in p.runs:
                    run.bold = True
        # 只把显式 th 的首行设为重复标题，避免把正文行误作表头。
        if r == 0 and cell['header']:
            pr = table.rows[0]._tr.get_or_add_trPr()
            if pr.find(qn('w:tblHeader')) is None:
                pr.append(OxmlElement('w:tblHeader'))


def build(html, meta, imgs, out):
    """meta 需含 cite/doi/author_line/orgn_line，缺哪项就少哪行"""
    soup = BeautifulSoup(html, 'html.parser')
    trad = is_trad(soup.get_text()[:20000])
    preclean(soup, trad)
    if meta.get('cite'):
        # 来源行整条跟随正文用字：知网元数据一律简体，《中國經學》这类繁體刊物
        # 直接照抄就成了简体刊名夹在繁體正文上头。刊名的正体写法以正文为准。
        meta['cite'] = _cc(meta['cite'], trad)
    doc = Document()
    st = doc.styles['Normal']; st.font.name = FONT; st.font.size = SZ_BODY
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.17)
    for name, size, mid in (('Heading 1', SZ_H0, True), ('Heading 2', SZ_H1, True),
                            ('Heading 3', SZ_H2, False), ('Heading 4', SZ_H3, False),
                            ('Heading 5', SZ_H4, False)):
        h = doc.styles[name]
        h.font.name = FONT; h.font.size = size; h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        h.paragraph_format.line_spacing = 1.5
        h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(6)
        if mid:
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    FN.set_continuous(doc)

    refs = references(soup)
    # 先取作者行：parse 会把 li 里的节点搬进包装 span，之后 ul 就取不到文字了
    crowd = {re.sub(r'\s+', '', t) for t in people(soup)}
    blocks, notes, source_images = parse(soup, refs)
    title = abstract = keywords = fund = bio = None
    body, byline, dropped, in_head = [], [], [], True
    for kind, nodes in blocks:
        if kind == 'table':
            body.append((kind, nodes)); in_head = False; continue
        txt = text_nodes(nodes)
        if kind == 'h1' and title is None:
            title = nodes; continue
        if kind == 'h1' and _is_en(txt):
            dropped.append(txt); continue   # 英文题名。用占比判定不用首字母——有的英文题名以中文引号起首
        if re.match(r'^摘\s*要[:：]', txt) and abstract is None:
            abstract = nodes; continue
        if re.match(r'^(关键词|關鍵詞)[:：]', txt) and keywords is None:
            keywords = nodes; continue
        if in_head and re.match(r'^作者[:：]', txt) and bio is None:
            # 各刊自带的作者简介：去掉标签，挂到作者行下做脚注；刊物没给就不出注
            bio = re.sub(r'^作者[:：]\s*', '', txt).strip(); continue
        if in_head and re.match(r'^基金[:：]', txt) and all(k == 't' for k, v in nodes):
            fund = re.sub(r'^基金[:：]\s*', '', txt).replace('；', '、').rstrip('。')
            fund = f"本文是{fund}。"; continue
        if in_head:
            if kind in ('h1', 'h2', 'h3', 'h4'):
                in_head = False
            elif re.sub(r'\s+', '', txt) in crowd:   # ul.Chapter-people：中文作者与单位上提到题名下
                (byline if _has_han(txt) else dropped).append(txt.strip()); continue
            elif _is_en(txt):
                dropped.append(txt); continue         # 英文摘要等首部西文
        body.append((kind, nodes))

    def center(size, before=0, after=0):
        return _newp(doc, size=size, spacing=1.5, align=WD_ALIGN_PARAGRAPH.CENTER,
                     before=before, after=after)

    if meta.get('cite'):
        _cnfont(center(SZ_META).add_run(meta['cite']), SZ_META)
    if meta.get('doi'):
        _cnfont(center(SZ_META, after=2).add_run("DOI：" + meta['doi']), SZ_META)
    if title:
        p = doc.add_paragraph(style='Heading 1'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _emit(p, title, SZ_H0, imgs)
        for r in p.runs:
            _cnfont(r, SZ_H0, bold=True)
        if fund:
            FN.add_ref(p, 1, sz_half=int(SZ_H0.pt * 2), mark='*')
    # 作者行统一用知网元数据合成的「姓名（單位）」，各刊自排的那两行只做兜底
    lines = [l for l in (meta.get('author_line'), meta.get('orgn_line')) if l]
    if lines:
        dropped.extend(byline)          # 被元数据取代的原作者行，验收要放行
    else:
        lines = byline
    if bio and not lines:               # 没有作者行可挂，简介退回首部段落，不留孤注
        body.insert(0, ('p', [('t', '作者：' + bio)])); bio = None
    for i, line in enumerate(lines):
        p = center(SZ_META)
        _cnfont(p.add_run(line), SZ_META)
        if i == 0 and bio:      # 简介注在作者名下，与题名下的基金星号注区分开
            FN.add_ref(p, 2, sz_half=int(SZ_META.pt * 2), mark='**' if fund else '*')
    last_head = None
    for part in (abstract, keywords):
        if part:
            last_head = _newp(doc, size=SZ_BODY, spacing=1.5, indent_chars=2)
            _emit(last_head, part, SZ_BODY, imgs)
    if last_head:                       # 摘要关键词与正文之间留一行余白
        last_head.paragraph_format.space_after = Pt(18)

    levels = {'h1': ('Heading 2', SZ_H1), 'h2': ('Heading 3', SZ_H2),
              'h3': ('Heading 4', SZ_H3), 'h4': ('Heading 5', SZ_H4),
              'h5': ('Heading 5', SZ_H4), 'h6': ('Heading 5', SZ_H4)}
    for kind, nodes in body:
        if kind == 'table':
            _emit_table(doc, nodes, imgs); continue
        txt = text_nodes(nodes)
        if kind == 'fig':                      # 表格／插图：独立成段，按版心宽度缩放
            idx = nodes[0][1]
            data = imgs.get(idx)
            p = _newp(doc, size=SZ_BODY, spacing=1.0, align=WD_ALIGN_PARAGRAPH.CENTER,
                      before=6, after=6)
            if data:
                try:
                    im = Image.open(io.BytesIO(data))
                    w = min(Cm(14.6), Pt(im.width * 0.75))     # 版心宽 14.66cm 封顶
                    p.add_run().add_picture(io.BytesIO(data), width=w,
                                            height=int(w * im.height / im.width))
                except Exception:
                    _cnfont(p.add_run("［图］"), SZ_BODY)
            else:
                _cnfont(p.add_run("［图］"), SZ_BODY)
            continue
        if kind == 'caption' or (kind in ('h4', 'h3') and re.match(r'^[表圖图]\s*\d+', txt)):
            p = _newp(doc, size=SZ_AUX, spacing=1.0,
                      align=WD_ALIGN_PARAGRAPH.CENTER, before=6)   # 表题／图题，不入大纲
            _emit(p, nodes, SZ_AUX, imgs)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.bold = True
            continue
        if kind == 'h1' and re.match(r'^第[一二三四五六七八九十]，', txt):
            # 知网把段内小标题（「第二，……」）标成了章节，降级为加粗正文而不是大标题。
            # 只对原本是标题的降级项加粗——普通段落哪怕也以「第二，」开头，也不能动
            p = _newp(doc, size=SZ_BODY, spacing=1.5, indent_chars=2)
            _emit(p, nodes, SZ_BODY, imgs)
            for r in p.runs:
                r.font.bold = True
            continue
        if kind in levels:
            style, size = levels[kind]
            p = doc.add_paragraph(style=style); _emit(p, nodes, size, imgs)
            for r in p.runs:
                _cnfont(r, size, bold=True)
        else:
            p = _newp(doc, size=SZ_BODY, spacing=1.5, indent_chars=2)
            _emit(p, nodes, SZ_BODY, imgs)

    if refs:
        p = doc.add_paragraph(style='Heading 2')
        _cnfont(p.add_run("參考文獻" if trad else "参考文献"), SZ_H1, bold=True)
        for r in refs:
            q = _newp(doc, size=SZ_AUX, spacing=1.5)
            q.paragraph_format.left_indent = Pt(28)
            q.paragraph_format.first_line_indent = Pt(-28)
            _cnfont(q.add_run(r), SZ_AUX)

    # 最终名称只指向完整 ZIP；先写同目录暂存文件、注入脚注并验收，再原子替换。
    from pathlib import Path
    out = Path(out); out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=out.parent, prefix='.' + out.stem,
                                     suffix='.docx', delete=False) as f:
        pending = Path(f.name)
    doc.save(pending)
    used = {int(x.get(qn('w:id'))) for x in doc.element.iter(qn('w:footnoteReference'))}
    fns = ([(1, fund, '*')] if fund else [])
    if bio:
        fns.append((2, bio, '**' if fund else '*'))
    fns += [(i + 3, n, None) for i, n in enumerate(notes) if i + 3 in used]
    FN.inject(pending, fns, sz_half=int(SZ_NOTE.pt * 2))
    report = QA.inspect_docx(pending, html, dropped=dropped)
    report.update(build_version=BUILD_VERSION, dropped=dropped,
                  missing_image_indices=sorted(set(range(len(source_images))) - set(imgs)))
    if report['status'] == 'invalid':
        REC.record(out, meta, report=report)
        raise ValueError('DOCX 未通过结构验收，原文件保留：' + str(pending))
    os.replace(pending, out)
    report['sha256'] = QA.fingerprint(out)
    REC.record(out, meta, report=report)
    return {"paras": len(body), "notes": len(fns) - bool(fund) - bool(bio), "fund": bool(fund),
            "imgs": report['docx']['images'], "trad": trad, "refs": len(refs),
            'quality': report, 'tables': report['docx']['tables']}
